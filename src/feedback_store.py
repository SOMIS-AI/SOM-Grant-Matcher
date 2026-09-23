"""
Faculty feedback verdicts — the 👍 / 👎 clicks from the digest emails
=====================================================================
Every match row in a personalised digest carries "Good match" / "Not relevant"
links into a Microsoft Form. The Form's export (one row per click) is dropped
into a local folder and imported here into `seed_data/feedback_verdicts.json`,
which ships with the container image like the Eval App keywords do.

The matcher reads the store at the start of every run (see
`load_feedback_index`) and uses it two ways:

  1. SUPPRESSION — a (faculty, grant) pair the faculty member rejected is never
     delivered to them again, nor is a re-post of the same call.
  2. PER-PERSON KEYWORD WEIGHTS — the keywords that anchored a rejected
     keyword match are down-weighted for that person only (a fact about them,
     not about the vocabulary); a "Good match" lifts them.

Only `self` verdicts (the faculty member rating their own match) feed the
matcher by default. `digest` verdicts — an admin rating someone else's match
from the shared digest — are kept in the store for review but not applied.

Match record format (built by emailer._match_feedback_id):
    email|grant#|run_date|confidence|match_type|rater
    email|ALL|run_date|-|-|self          "none of these are relevant"
    email|OPTOUT|run_date|-|-|self       footer opt-out link
"""
from __future__ import annotations

import json
import logging
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

DEFAULT_STORE_PATH = Path("seed_data/feedback_verdicts.json")

VERDICT_GOOD = "good"
VERDICT_NOT_RELEVANT = "not_relevant"
VERDICT_ALL_NOT_RELEVANT = "all_not_relevant"
VERDICT_OPTOUT = "optout"


# ── Store I/O ────────────────────────────────────────────────────────────────

def load_store(path: Path = DEFAULT_STORE_PATH) -> dict:
    path = Path(path)
    if not path.exists():
        return {"updated_at": None, "sources": [], "verdicts": []}
    try:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        data.setdefault("sources", [])
        data.setdefault("verdicts", [])
        return data
    except (OSError, json.JSONDecodeError) as e:
        logger.warning(f"feedback store unreadable ({e}); starting empty")
        return {"updated_at": None, "sources": [], "verdicts": []}


def save_store(store: dict, path: Path = DEFAULT_STORE_PATH) -> None:
    from atomic_io import atomic_write_json
    store["updated_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(path, store, indent=1)


# ── Parsing the Form export ──────────────────────────────────────────────────

def _norm_title(title: str) -> str:
    return re.sub(r"[^a-z0-9 ]+", " ", (title or "").lower()).strip()


def _verdict_from_text(text: str) -> Optional[str]:
    t = (text or "").strip().lower()
    if t.startswith("good"):
        return VERDICT_GOOD
    if t.startswith("not relevant"):
        return VERDICT_NOT_RELEVANT
    return None


def parse_match_record(record: str) -> dict:
    """email|grant|run_date|conf|match_type|rater -> dict (missing parts blank)."""
    parts = [p.strip() for p in str(record or "").split("|")]
    parts += [""] * (6 - len(parts))
    email, grant, run_date, conf, match_type, rater = parts[:6]
    try:
        conf_i: Optional[int] = int(conf)
    except ValueError:
        conf_i = None
    return {
        "email": email.lower(),
        "grant": grant,
        "run_date": run_date,
        "confidence": conf_i,
        "match_type": match_type.lower(),
        "rater": (rater or "self").lower(),
    }


def _archive_lookup(archive_dir: Optional[Path]) -> dict:
    """(grant_number, email) -> {title, keywords, match_type} from the Daily
    match workbooks, so a verdict can carry the keywords that anchored the
    match. Local-only convenience: on the server the archive is not present
    and the importer simply stores what the Form record carries."""
    idx: dict = {}
    if not archive_dir:
        return idx
    try:
        import openpyxl
    except ImportError:
        return idx
    for path in sorted(Path(archive_dir).glob("UMSOM_Grant_Matches_Daily_*.xlsx")):
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
        except Exception as e:
            logger.warning(f"archive workbook skipped {path.name}: {e}")
            continue
        for name in wb.sheetnames[1:]:
            rows = list(wb[name].iter_rows(values_only=True))
            if not rows:
                continue
            title = str(rows[0][0] or "")
            gnum = next((str(r[1]) for r in rows[:8]
                         if r and str(r[0] or "").startswith("Grant Number")), "")
            hi = next((i for i, r in enumerate(rows) if r and str(r[0]) == "Faculty"), None)
            if hi is None:
                continue
            hdr = [str(c) for c in rows[hi]]
            for r in rows[hi + 1:]:
                if not r or not r[0]:
                    continue
                d = dict(zip(hdr, r))
                em = str(d.get("Email") or "").strip().lower()
                if not em:
                    continue
                kws = [k.strip() for k in str(d.get("Matched Keywords") or "").split(",")
                       if k.strip() and not k.strip().startswith("≈")]
                idx.setdefault((gnum, em), {"title": title, "keywords": kws,
                                            "match_type": str(d.get("Match Type") or "")})
    return idx


def import_export(xlsx_path: Path, store: dict, archive_dir: Optional[Path] = None) -> dict:
    """Merge one Form export into `store`. Rows are keyed by the Form's Id so a
    re-imported export never duplicates a verdict. Returns stats."""
    import openpyxl
    stats = {"rows_total": 0, "added": 0, "duplicate": 0, "unparseable": 0,
             "keywords_resolved": 0}
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    if not rows:
        return stats
    hdr = [str(c) for c in rows[0]]
    col = {h: i for i, h in enumerate(hdr)}
    need = ("Id", "Match record", "Matching Feedback")
    if any(k not in col for k in need):
        raise ValueError(f"unexpected columns {hdr}; need {need}")
    comment_col = next((i for h, i in col.items() if h.lower().startswith("any additional")), None)
    time_col = col.get("Start time")

    known = {(v.get("source_file"), v.get("form_id")) for v in store["verdicts"]}
    archive = _archive_lookup(archive_dir)
    for r in rows[1:]:
        if not r or r[col["Id"]] is None:
            continue
        stats["rows_total"] += 1
        form_id = str(r[col["Id"]])
        if (xlsx_path.name, form_id) in known:
            stats["duplicate"] += 1
            continue
        rec = parse_match_record(r[col["Match record"]])
        grant = rec["grant"]
        if grant.upper() == "OPTOUT":
            verdict = VERDICT_OPTOUT
        elif grant.upper() == "ALL":
            verdict = VERDICT_ALL_NOT_RELEVANT
        else:
            verdict = _verdict_from_text(str(r[col["Matching Feedback"]] or ""))
        if not verdict or not rec["email"]:
            stats["unparseable"] += 1
            continue
        entry = {
            "form_id": form_id,
            "source_file": xlsx_path.name,
            "responded_at": str(r[time_col])[:19] if time_col is not None and r[time_col] else "",
            "verdict": verdict,
            **rec,
            "comment": str(r[comment_col] or "").strip() if comment_col is not None else "",
            "grant_title": "",
            "keywords": [],
        }
        hit = archive.get((grant, rec["email"]))
        if hit:
            entry["grant_title"] = hit["title"]
            entry["keywords"] = hit["keywords"]
            stats["keywords_resolved"] += 1
        store["verdicts"].append(entry)
        known.add((xlsx_path.name, form_id))
        stats["added"] += 1
    store["sources"].append({
        "filename": xlsx_path.name,
        "imported_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "stats": stats,
    })
    return stats


# ── What the matcher consumes ────────────────────────────────────────────────

def build_feedback_index(store: dict, *, use_third_party: bool = False,
                         rejected_multiplier: float = 0.5,
                         good_multiplier: float = 1.2) -> dict:
    """
    Turn the verdict list into the structures the matcher needs:

      suppress        {(email, grant_number)}       never deliver again
      suppress_titles {(email, normalised title)}   same call re-posted
      kw_weights      {email: {kw_norm: multiplier}} per-person keyword weights
      optout          {email}                        asked to stop receiving
      counts          summary for the diagnostic
    """
    from matcher import normalize  # sibling module; matcher imports us lazily

    suppress, suppress_titles, optout = set(), set(), set()
    kw_weights: dict = {}
    counts = {"verdicts": len(store.get("verdicts", [])), "applied": 0,
              "third_party_skipped": 0, "good": 0, "not_relevant": 0, "optout": 0}
    for v in store.get("verdicts", []):
        email = (v.get("email") or "").lower()
        if not email or email.startswith("name:"):
            continue
        if v.get("rater", "self") != "self" and not use_third_party:
            counts["third_party_skipped"] += 1
            continue
        verdict = v.get("verdict")
        if verdict == VERDICT_OPTOUT:
            optout.add(email)
            counts["optout"] += 1
            counts["applied"] += 1
            continue
        if verdict == VERDICT_ALL_NOT_RELEVANT:
            # We do not know which grants were in that digest here; the
            # run-date-wide rejection is recorded for review only.
            continue
        grant = v.get("grant") or ""
        if not grant:
            continue
        counts["applied"] += 1
        if verdict == VERDICT_NOT_RELEVANT:
            counts["not_relevant"] += 1
            suppress.add((email, grant))
            if v.get("grant_title"):
                suppress_titles.add((email, _norm_title(v["grant_title"])))
            mult = rejected_multiplier
        elif verdict == VERDICT_GOOD:
            counts["good"] += 1
            mult = good_multiplier
        else:
            continue
        # keyword weights only for lexically grounded matches
        if v.get("match_type") in ("keyword", "both"):
            w = kw_weights.setdefault(email, {})
            for kw in v.get("keywords") or []:
                k = normalize(kw).strip()
                if k:
                    w[k] = round(w.get(k, 1.0) * mult, 4)
    return {"suppress": suppress, "suppress_titles": suppress_titles,
            "kw_weights": kw_weights, "optout": optout, "counts": counts,
            "norm_title": _norm_title}


def load_feedback_index(config: Optional[dict] = None,
                        path: Path = DEFAULT_STORE_PATH) -> Optional[dict]:
    """Index from the store on disk, honouring matching.feedback config.
    Returns None when disabled or the store is absent/empty."""
    fb_cfg = ((config or {}).get("matching", {}) or {}).get("feedback", {}) or {}
    if not fb_cfg.get("enabled", True):
        return None
    store_path = Path(fb_cfg.get("store_path") or path)
    store = load_store(store_path)
    if not store.get("verdicts"):
        return None
    return build_feedback_index(
        store,
        use_third_party=bool(fb_cfg.get("use_third_party_verdicts", False)),
        rejected_multiplier=float(fb_cfg.get("rejected_keyword_multiplier", 0.5)),
        good_multiplier=float(fb_cfg.get("good_keyword_multiplier", 1.2)),
    )


# ── CLI ──────────────────────────────────────────────────────────────────────

def _cli(argv: list[str]) -> int:
    """python -m src.feedback_store <form-export.xlsx> [--archive <Diag Files dir>]

    Merges the export into seed_data/feedback_verdicts.json. With --archive,
    each verdict is enriched with the grant title and matched keywords from
    the Daily match workbooks so keyword weights can be derived."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    args = [a for a in argv[1:]]
    archive = None
    if "--archive" in args:
        i = args.index("--archive")
        archive = Path(args[i + 1])
        del args[i:i + 2]
    if not args:
        print("Usage: python -m src.feedback_store <xlsx> [<xlsx>...] [--archive <dir>]")
        return 2
    store = load_store()
    for arg in args:
        p = Path(arg)
        if not p.exists():
            print(f"  ✗ Missing: {p}")
            continue
        print(f"  Importing {p.name}...")
        try:
            s = import_export(p, store, archive_dir=archive)
        except Exception as e:
            print(f"  ✗ {p.name}: {type(e).__name__}: {e}")
            print("  ABORTED — the store on disk was NOT modified.")
            return 1
        print(f"    rows={s['rows_total']} added={s['added']} duplicate={s['duplicate']} "
              f"unparseable={s['unparseable']} keywords_resolved={s['keywords_resolved']}")
    save_store(store)
    idx = build_feedback_index(store)
    c = idx["counts"]
    print(f"\n[OK] Store has {c['verdicts']} verdicts; matcher will apply {c['applied']} "
          f"(good {c['good']}, not relevant {c['not_relevant']}, opt-out {c['optout']}; "
          f"{c['third_party_skipped']} third-party skipped)")
    print(f"  -> {DEFAULT_STORE_PATH}")
    return 0


if __name__ == "__main__":
    if str(Path(__file__).parent) not in sys.path:
        sys.path.insert(0, str(Path(__file__).parent))
    raise SystemExit(_cli(sys.argv))
